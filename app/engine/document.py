"""文档解析（§1.3.5 合同解析字段要求 —— 文本获取与定位信息）。

支持三类合同文件：
  - .docx          直接抽取段落文本（页码统一记为第 1 页）
  - 文本型 PDF     逐页抽取文本层，段落按行切分
  - 扫描型 PDF     无文本层，逐页渲染为图像后交由 OCR 引擎识别

解析失败时抛出 `DocumentParseError` 并携带明确的失败原因，
对应 §1.3.4 的两类阻塞：`文档内容为空` / `图片无法识别`。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pymupdf

from .ocr import get_ocr_engine

# OCR 渲染分辨率
OCR_DPI = 150

# 判定「有文本层」的最小字符数：低于该值视为扫描件
TEXT_LAYER_MIN_CHARS = 20


class DocumentParseError(Exception):
    """文档解析失败，message 即阻塞原因（§1.3.4）。"""

    def __init__(self, message: str, detail: str = "") -> None:
        super().__init__(message)
        self.message = message
        self.detail = detail


@dataclass
class Paragraph:
    """段落及其位置信息。"""

    page_no: int
    index: int          # 页内段落序号，从 1 开始
    text: str

    @property
    def position(self) -> str:
        return f"第{self.page_no}页/第{self.index}段"


@dataclass
class ParsedDocument:
    """解析后的文档结构。"""

    file_name: str
    source_path: str
    parse_type: str                     # docx / pdf_text / pdf_ocr
    page_count: int
    paragraphs: list[Paragraph] = field(default_factory=list)
    ocr_engine: str = ""
    warnings: list[str] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n".join(item.text for item in self.paragraphs)

    @property
    def char_count(self) -> int:
        return sum(len(item.text) for item in self.paragraphs)


def _split_lines(text: str, page_no: int, start_index: int = 1) -> list[Paragraph]:
    """把一页文本按行切成段落，保留页内序号。"""
    paragraphs: list[Paragraph] = []
    index = start_index
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        paragraphs.append(Paragraph(page_no=page_no, index=index, text=line))
        index += 1
    return paragraphs


def _parse_docx(path: Path) -> ParsedDocument:
    from docx import Document

    try:
        document = Document(str(path))
    except Exception as exc:  # 损坏的 docx
        raise DocumentParseError("文档内容为空", f"docx 解析失败：{exc}") from exc

    paragraphs = [
        Paragraph(page_no=1, index=index, text=item.text.strip())
        for index, item in enumerate(
            (p for p in document.paragraphs if p.text and p.text.strip()), start=1
        )
    ]

    if not paragraphs:
        raise DocumentParseError("文档内容为空", "docx 文档不包含任何有效段落文本")

    return ParsedDocument(
        file_name=path.name,
        source_path=str(path),
        parse_type="docx",
        page_count=1,
        paragraphs=paragraphs,
    )


def _parse_pdf(path: Path) -> ParsedDocument:
    try:
        doc = pymupdf.open(str(path))
    except Exception as exc:
        raise DocumentParseError("文档内容为空", f"PDF 打开失败：{exc}") from exc

    try:
        if doc.page_count == 0:
            raise DocumentParseError("文档内容为空", "PDF 不包含任何页面")

        page_texts: list[str] = []
        for page in doc:
            try:
                page_texts.append(page.get_text("text") or "")
            except Exception:
                page_texts.append("")

        total_chars = sum(len(text.strip()) for text in page_texts)

        # 有文本层：直接按文本解析
        if total_chars >= TEXT_LAYER_MIN_CHARS:
            paragraphs: list[Paragraph] = []
            for page_no, text in enumerate(page_texts, start=1):
                paragraphs.extend(_split_lines(text, page_no))
            if not paragraphs:
                raise DocumentParseError("文档内容为空", "PDF 文本层为空")
            return ParsedDocument(
                file_name=path.name,
                source_path=str(path),
                parse_type="pdf_text",
                page_count=doc.page_count,
                paragraphs=paragraphs,
            )

        # 无文本层：判定为扫描件，走 OCR
        return _parse_scanned_pdf(doc, path)
    finally:
        doc.close()


def _parse_scanned_pdf(doc: "pymupdf.Document", path: Path) -> ParsedDocument:
    engine = get_ocr_engine()
    slug = path.stem
    paragraphs: list[Paragraph] = []
    warnings: list[str] = []
    recognized_pages = 0

    for page_no, page in enumerate(doc, start=1):
        try:
            pixmap = page.get_pixmap(dpi=OCR_DPI)
            image = pixmap.tobytes("png")
        except Exception as exc:
            warnings.append(f"第{page_no}页图像渲染失败：{exc}")
            continue

        try:
            text = engine.recognize(image, page_no, slug)
        except Exception as exc:
            warnings.append(f"第{page_no}页 OCR 调用失败：{exc}")
            continue

        if not text.strip():
            warnings.append(f"第{page_no}页未识别出文本")
            continue

        recognized_pages += 1
        paragraphs.extend(_split_lines(text, page_no))

    if not paragraphs:
        raise DocumentParseError(
            "图片无法识别",
            f"扫描件共 {doc.page_count} 页，OCR 引擎（{engine.describe()}）未识别出任何文本",
        )

    return ParsedDocument(
        file_name=path.name,
        source_path=str(path),
        parse_type="pdf_ocr",
        page_count=doc.page_count,
        paragraphs=paragraphs,
        ocr_engine=engine.describe(),
        warnings=warnings,
    )


def parse_document(path: str | Path) -> ParsedDocument:
    """解析合同文档，返回带位置信息的段落结构。

    失败时抛出 DocumentParseError，message 为阻塞原因。
    """
    target = Path(path)
    if not target.is_file():
        raise DocumentParseError("附件缺失", f"待解析文件不存在：{target}")

    suffix = target.suffix.lower()
    if suffix == ".docx":
        return _parse_docx(target)
    if suffix == ".pdf":
        return _parse_pdf(target)

    raise DocumentParseError("文档内容为空", f"暂不支持的合同文件格式：{suffix or '(无扩展名)'}")
